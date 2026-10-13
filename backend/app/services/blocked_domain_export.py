"""The Blocked page's "Domains" list as a downloadable file -- every row, not
one page. Built in memory like actor_export.py: one row per domain, a few
thousand at most (blocked_domain_service.EXPORT_MAX_ROWS)."""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.schemas.blocked_domains import BlockedDomainRow
from app.services.export_service import _escape_csv_formula, _sanitize_xlsx_value

_TEXT = {
    "uz": {
        "headers": [
            "№", "Domen", "Bloklangan (403)", "Login so'rovi (407)", "Klientlar", "Filiallar",
            "Proksilar", "Birinchi (UTC)", "Oxirgi (UTC)",
        ],
        "sheet": "Domenlar", "info": "Ma'lumot",
        "branch": "Filial", "all_branches": "Barcha filiallar", "range": "Oraliq (UTC)",
        "proxy": "Proksi", "all_proxies": "Barcha proksilar", "kind": "Javob turi",
        "search": "Qidiruv", "sort": "Saralash", "rows": "Qatorlar", "created": "Yaratilgan (UTC)",
        "kinds": {"denied": "Haqiqiy blok", "auth": "Login so'rovi (407)", "all": "Barcha bloklar"},
    },
    "en": {
        "headers": [
            "#", "Domain", "Blocked (403)", "Login required (407)", "Clients", "Branches",
            "Proxies", "First (UTC)", "Last (UTC)",
        ],
        "sheet": "Domains", "info": "Info",
        "branch": "Branch", "all_branches": "All branches", "range": "Range (UTC)",
        "proxy": "Proxy", "all_proxies": "All proxies", "kind": "Response type",
        "search": "Search", "sort": "Sort", "rows": "Rows", "created": "Created (UTC)",
        "kinds": {"denied": "Real blocks", "auth": "Login required (407)", "all": "All blocks"},
    },
    "ru": {
        "headers": [
            "№", "Домен", "Заблокировано (403)", "Запрос логина (407)", "Клиенты", "Филиалы",
            "Прокси", "Первый (UTC)", "Последний (UTC)",
        ],
        "sheet": "Домены", "info": "Сведения",
        "branch": "Филиал", "all_branches": "Все филиалы", "range": "Период (UTC)",
        "proxy": "Прокси", "all_proxies": "Все прокси", "kind": "Тип ответа",
        "search": "Поиск", "sort": "Сортировка", "rows": "Строк", "created": "Создано (UTC)",
        "kinds": {"denied": "Реальные блокировки", "auth": "Запрос логина (407)", "all": "Все блокировки"},
    },
}


def _values(index: int, row: BlockedDomainRow) -> list[object]:
    return [
        index,
        _escape_csv_formula(row.domain),
        row.denied_count,
        row.auth_count,
        row.client_count,
        ", ".join(row.branches),
        ", ".join(row.proxies),
        f"{row.first_seen:%Y-%m-%d %H:%M}",
        f"{row.last_seen:%Y-%m-%d %H:%M}",
    ]


def domains_csv(rows: list[BlockedDomainRow], lang: str) -> bytes:
    text = _TEXT.get(lang, _TEXT["uz"])
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(text["headers"])
    for i, row in enumerate(rows, start=1):
        writer.writerow(_values(i, row))
    # BOM so Excel opens the UTF-8 text correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


def domain_names_csv(rows: list[BlockedDomainRow]) -> bytes:
    """Just the domains, one per line, no header -- ready to paste into a
    Squid ACL or a firewall list."""
    return "".join(f"{_escape_csv_formula(row.domain)}\n" for row in rows).encode("utf-8")


def domains_xlsx(
    rows: list[BlockedDomainRow],
    lang: str,
    *,
    branch: str | None,
    proxy: str | None,
    since: datetime,
    until: datetime,
    kind: str,
    search: str | None,
    sort: str,
) -> bytes:
    text = _TEXT.get(lang, _TEXT["uz"])
    wb = Workbook()
    ws = wb.active
    ws.title = text["sheet"]
    ws.append(text["headers"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=1):
        ws.append([_sanitize_xlsx_value(v) for v in _values(i, row)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in enumerate([7, 40, 16, 18, 11, 18, 26, 17, 17], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    info = wb.create_sheet(text["info"])
    for label, value in [
        (text["branch"], branch or text["all_branches"]),
        (text["proxy"], proxy or text["all_proxies"]),
        (text["range"], f"{since:%Y-%m-%d %H:%M} — {until:%Y-%m-%d %H:%M}"),
        (text["kind"], text["kinds"].get(kind, kind)),
        (text["search"], _escape_csv_formula(search or "—")),
        (text["sort"], sort),
        (text["rows"], len(rows)),
        (text["created"], f"{datetime.now(until.tzinfo):%Y-%m-%d %H:%M}"),
    ]:
        info.append([label, _sanitize_xlsx_value(value)])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 40

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
