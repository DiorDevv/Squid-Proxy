"""The Blocked page's Analysis list as a downloadable file -- every group,
not one page. Built in memory like actor_export.py: one row per group, a
few thousand at most (blocked_analysis_service.EXPORT_MAX_ROWS)."""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.branch_scope import BranchScope, scope_label
from app.schemas.blocked_domains import BlockGroupRow
from app.services.export_service import _escape_csv_formula, _sanitize_xlsx_value

_TEXT = {
    "uz": {
        "groups": {
            "domain": "Domen", "client": "Klient IP", "user": "Login", "subnet": "Subnet",
            "branch": "Filial", "proxy": "Proksi",
        },
        "headers": [
            "Bloklangan (403)", "Login so'rovi (407)", "Klientlar", "Loginlar", "Domenlar",
            "Filiallar", "Proksilar", "Birinchi (UTC)", "Oxirgi (UTC)",
        ],
        "none": "(yo'q)", "info": "Ma'lumot", "filters": "Filtrlar",
        "branch": "Filial", "all_branches": "Barcha filiallar", "range": "Oraliq (UTC)",
        "kind": "Javob turi", "group": "Guruhlash", "sort": "Saralash", "rows": "Qatorlar",
        "created": "Yaratilgan (UTC)",
        "kinds": {"denied": "Haqiqiy blok (403)", "auth": "Login so'rovi (407)", "all": "Barcha bloklar"},
    },
    "en": {
        "groups": {
            "domain": "Domain", "client": "Client IP", "user": "Login", "subnet": "Subnet",
            "branch": "Branch", "proxy": "Proxy",
        },
        "headers": [
            "Blocked (403)", "Login required (407)", "Clients", "Logins", "Domains",
            "Branches", "Proxies", "First (UTC)", "Last (UTC)",
        ],
        "none": "(none)", "info": "Info", "filters": "Filters",
        "branch": "Branch", "all_branches": "All branches", "range": "Range (UTC)",
        "kind": "Response type", "group": "Grouped by", "sort": "Sort", "rows": "Rows",
        "created": "Created (UTC)",
        "kinds": {"denied": "Real blocks (403)", "auth": "Login required (407)", "all": "All blocks"},
    },
    "ru": {
        "groups": {
            "domain": "Домен", "client": "IP клиента", "user": "Логин", "subnet": "Подсеть",
            "branch": "Филиал", "proxy": "Прокси",
        },
        "headers": [
            "Заблокировано (403)", "Запрос логина (407)", "Клиенты", "Логины", "Домены",
            "Филиалы", "Прокси", "Первый (UTC)", "Последний (UTC)",
        ],
        "none": "(нет)", "info": "Сведения", "filters": "Фильтры",
        "branch": "Филиал", "all_branches": "Все филиалы", "range": "Период (UTC)",
        "kind": "Тип ответа", "group": "Группировка", "sort": "Сортировка", "rows": "Строк",
        "created": "Создано (UTC)",
        "kinds": {
            "denied": "Реальные блокировки (403)", "auth": "Запрос логина (407)", "all": "Все блокировки",
        },
    },
}


def _text(lang: str) -> dict:
    return _TEXT.get(lang, _TEXT["uz"])


def _values(index: int, row: BlockGroupRow, none: str) -> list[object]:
    return [
        index,
        _escape_csv_formula(row.key or none),
        row.denied_count,
        row.auth_count,
        row.client_count,
        row.user_count,
        row.domain_count,
        ", ".join(row.branches),
        ", ".join(row.proxies),
        f"{row.first_seen:%Y-%m-%d %H:%M}",
        f"{row.last_seen:%Y-%m-%d %H:%M}",
    ]


def groups_csv(rows: list[BlockGroupRow], group: str, lang: str) -> bytes:
    text = _text(lang)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["№", text["groups"][group], *text["headers"]])
    for i, row in enumerate(rows, start=1):
        writer.writerow(_values(i, row, text["none"]))
    # BOM so Excel opens the UTF-8 text correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


def group_keys_csv(rows: list[BlockGroupRow]) -> bytes:
    """Just the keys (domains / IPs / logins ...), one per line, no header --
    ready to paste into a Squid ACL or a firewall list."""
    return "".join(f"{_escape_csv_formula(row.key)}\n" for row in rows if row.key).encode("utf-8")


def groups_xlsx(
    rows: list[BlockGroupRow],
    group: str,
    lang: str,
    *,
    branch: BranchScope,
    since: datetime,
    until: datetime,
    kind: str,
    sort: str,
    filters: list[tuple[str, str]],
) -> bytes:
    text = _text(lang)
    wb = Workbook()
    ws = wb.active
    ws.title = text["groups"][group]
    ws.append(["№", text["groups"][group], *text["headers"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=1):
        ws.append([_sanitize_xlsx_value(v) for v in _values(i, row, text["none"])])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in enumerate([7, 36, 16, 18, 11, 11, 11, 18, 26, 17, 17], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    info = wb.create_sheet(text["info"])
    for label, value in [
        (text["group"], text["groups"][group]),
        (text["branch"], scope_label(branch) or text["all_branches"]),
        (text["range"], f"{since:%Y-%m-%d %H:%M} — {until:%Y-%m-%d %H:%M}"),
        (text["kind"], text["kinds"].get(kind, kind)),
        *[(name, _escape_csv_formula(value)) for name, value in filters],
        (text["sort"], sort),
        (text["rows"], len(rows)),
        (text["created"], f"{datetime.now(until.tzinfo):%Y-%m-%d %H:%M}"),
    ]:
        info.append([label, _sanitize_xlsx_value(value)])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
