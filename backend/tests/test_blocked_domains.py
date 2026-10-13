"""The Blocked page's "Domains" view: the hourly aggregates the aggregator
fills (403 vs 407, per proxy and client), the grouped list built on them,
its export, and the events list's block_kind filter."""

import dataclasses
import io
from datetime import UTC, datetime, timedelta

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.models.audit_log import AuditAction, AuditLogEntry
from app.models.blocked_domain_aggregate import (
    BlockedDomainClientHourlyAggregate,
    BlockedDomainHourlyAggregate,
)
from app.models.raw_event import RawEvent
from app.services import blocked_domain_service
from app.services.aggregator import Aggregator
from app.services.event_store import RingBuffer
from app.services.log_parser import parse_line

HOUR = datetime.now(UTC).replace(minute=0, second=0, microsecond=0) - timedelta(hours=2)
NOW = datetime.now(UTC)
SINCE = NOW - timedelta(hours=24)


def _line(domain: str, client_ip: str, status: str) -> str:
    return (
        f"1737100800.123 45 {client_ip} {status} 0 GET http://{domain}/ - HIER_NONE/- -"
    )


async def test_flush_splits_denied_and_login_per_hour_proxy_and_client(db_engine, monkeypatch):
    session_factory = async_sessionmaker(db_engine, expire_on_commit=False, class_=AsyncSession)
    import app.services.aggregator as aggregator_module

    monkeypatch.setattr(aggregator_module, "AsyncSessionLocal", session_factory)

    ring_buffer = RingBuffer(max_events=100)
    via77 = lambda line: dataclasses.replace(parse_line(line), proxy="172.25.42.77")  # noqa: E731
    ring_buffer.append(via77(_line("ctldl.windowsupdate.com", "10.0.0.1", "TCP_DENIED/403")))
    ring_buffer.append(via77(_line("ctldl.windowsupdate.com", "10.0.0.1", "TCP_DENIED/403")))
    ring_buffer.append(via77(_line("ctldl.windowsupdate.com", "10.0.0.2", "TCP_DENIED/407")))
    ring_buffer.append(parse_line(_line("ctldl.windowsupdate.com", "10.0.0.3", "TCP_DENIED/403")))
    ring_buffer.append(parse_line(_line("allowed.example", "10.0.0.1", "TCP_MISS/200")))
    aggregator = Aggregator(ring_buffer=ring_buffer, interval_seconds=60)
    await aggregator.flush()
    # A second flush into the same hour adds to the rows, not replaces them.
    ring_buffer.append(via77(_line("ctldl.windowsupdate.com", "10.0.0.1", "TCP_DENIED/403")))
    await aggregator.flush()

    async with session_factory() as session:
        domain_rows = {
            row.proxy: row
            for row in (await session.execute(select(BlockedDomainHourlyAggregate))).scalars()
        }
        assert set(domain_rows) == {"172.25.42.77", ""}  # unlabeled source stored as ""
        assert (domain_rows["172.25.42.77"].denied_count, domain_rows["172.25.42.77"].auth_count) == (3, 1)
        assert (domain_rows[""].denied_count, domain_rows[""].auth_count) == (1, 0)
        assert domain_rows[""].bucket_ts.replace(tzinfo=UTC) == datetime(2025, 1, 17, 8, tzinfo=UTC)

        client_rows = {
            row.client_ip: (row.denied_count, row.auth_count)
            for row in (await session.execute(select(BlockedDomainClientHourlyAggregate))).scalars()
        }
        assert client_rows == {"10.0.0.1": (3, 0), "10.0.0.2": (0, 1), "10.0.0.3": (1, 0)}


def _agg(domain, denied, auth, *, branch="main", proxy="", hour=HOUR):
    return BlockedDomainHourlyAggregate(
        bucket_ts=hour, branch=branch, proxy=proxy, domain=domain, denied_count=denied, auth_count=auth
    )


def _client(domain, ip, denied, auth, *, branch="main", proxy="", hour=HOUR):
    return BlockedDomainClientHourlyAggregate(
        bucket_ts=hour, branch=branch, proxy=proxy, domain=domain, client_ip=ip,
        denied_count=denied, auth_count=auth,
    )


async def _seed(db: AsyncSession) -> None:
    db.add_all([
        _agg("crl.usertrust.com", 90, 10, proxy="172.25.42.77"),
        _agg("crl.usertrust.com", 10, 0, branch="filial", hour=HOUR - timedelta(hours=1)),
        _client("crl.usertrust.com", "10.0.0.1", 50, 0, proxy="172.25.42.77"),
        _client("crl.usertrust.com", "10.0.0.2", 40, 10, proxy="172.25.42.77"),
        _client("crl.usertrust.com", "10.0.0.3", 10, 0, branch="filial", hour=HOUR - timedelta(hours=1)),
        _agg("crl.sectigo.com", 60, 0, proxy="172.25.42.24"),
        _client("crl.sectigo.com", "10.0.0.1", 60, 0, proxy="172.25.42.24"),
        # login-only: not a block, so absent from the default list
        _agg("crl.apple.com", 0, 6, proxy="172.25.42.24"),
        _client("crl.apple.com", "10.0.0.9", 0, 6, proxy="172.25.42.24"),
        _agg("ctldl.windowsupdate.com", 500, 0, branch="filial"),
        _client("ctldl.windowsupdate.com", "10.0.0.4", 500, 0, branch="filial"),
        # outside the 24h range
        _agg("old.example", 999, 0, hour=HOUR - timedelta(days=3)),
    ])
    await db.commit()


async def _list(db, **kwargs):
    return await blocked_domain_service.get_blocked_domains(db, SINCE, NOW, **kwargs)


async def test_default_list_is_real_blocks_by_count(db_session: AsyncSession):
    await _seed(db_session)
    result = await _list(db_session)
    assert [r.domain for r in result.items] == ["ctldl.windowsupdate.com", "crl.usertrust.com", "crl.sectigo.com"]
    assert result.total == 3
    usertrust = result.items[1]
    assert (usertrust.denied_count, usertrust.auth_count, usertrust.client_count) == (100, 10, 3)
    assert usertrust.branches == ["filial", "main"]
    assert usertrust.proxies == ["172.25.42.77"]
    assert usertrust.first_seen < usertrust.last_seen
    assert (result.denied_total, result.auth_total) == (660, 16)
    assert result.available_proxies == ["172.25.42.24", "172.25.42.77"]


async def test_kind_search_proxy_and_paging(db_session: AsyncSession):
    await _seed(db_session)
    login = await _list(db_session, kind="auth")
    assert [(r.domain, r.client_count) for r in login.items] == [("crl.usertrust.com", 1), ("crl.apple.com", 1)]

    crl = await _list(db_session, search="crl.", kind="all")
    assert {r.domain for r in crl.items} == {"crl.usertrust.com", "crl.sectigo.com", "crl.apple.com"}

    via77 = await _list(db_session, proxy="172.25.42.77")
    assert [(r.domain, r.denied_count, r.client_count) for r in via77.items] == [("crl.usertrust.com", 90, 2)]
    # the proxy filter keeps offering every proxy
    assert via77.available_proxies == ["172.25.42.24", "172.25.42.77"]

    seen = []
    for offset in (0, 2):
        page = await _list(db_session, limit=2, offset=offset)
        assert page.total == 3
        seen += [r.domain for r in page.items]
    assert seen == ["ctldl.windowsupdate.com", "crl.usertrust.com", "crl.sectigo.com"]

    by_name = await _list(db_session, sort="domain")
    assert [r.domain for r in by_name.items] == sorted(r.domain for r in by_name.items)


async def test_branch_scope(db_session: AsyncSession):
    await _seed(db_session)
    result = await _list(db_session, branch="filial")
    assert [(r.domain, r.denied_count, r.client_count) for r in result.items] == [
        ("ctldl.windowsupdate.com", 500, 1),
        ("crl.usertrust.com", 10, 1),
    ]


async def test_route_lists_domains(app_client: AsyncClient, viewer_token, auth_headers, db_session):
    await _seed(db_session)
    response = await app_client.get(
        "/api/blocked-domains", params={"range": "24h", "search": "crl."}, headers=auth_headers(viewer_token)
    )
    assert response.status_code == 200
    body = response.json()
    assert [row["domain"] for row in body["items"]] == ["crl.usertrust.com", "crl.sectigo.com"]


async def test_export_domain_names_only(app_client: AsyncClient, admin_token, auth_headers, db_session):
    await _seed(db_session)
    response = await app_client.get(
        "/api/blocked-domains/export",
        params={"range": "24h", "format": "csv", "columns": "domain", "search": "crl."},
        headers=auth_headers(admin_token),
    )
    assert response.status_code == 200
    assert response.content.decode("utf-8").splitlines() == ["crl.usertrust.com", "crl.sectigo.com"]
    entries = (
        await db_session.execute(select(AuditLogEntry).where(AuditLogEntry.action == AuditAction.EXPORT_CREATED))
    ).scalars().all()
    assert [e.detail for e in entries] == [
        "blocked-domains format=csv columns=domain kind=denied rows=2 search=crl."
    ]


async def test_export_xlsx_has_every_row(app_client: AsyncClient, admin_token, auth_headers, db_session):
    await _seed(db_session)
    response = await app_client.get(
        "/api/blocked-domains/export",
        params={"range": "24h", "format": "xlsx", "lang": "uz"},
        headers=auth_headers(admin_token),
    )
    assert response.status_code == 200
    assert 'filename="bloklangan-domenlar-barcha-' in response.headers["content-disposition"]
    values = list(load_workbook(io.BytesIO(response.content))["Domenlar"].iter_rows(values_only=True))
    assert values[0][1] == "Domen"
    assert [v[1] for v in values[1:]] == ["ctldl.windowsupdate.com", "crl.usertrust.com", "crl.sectigo.com"]
    assert values[2][2:5] == (100, 10, 3)


async def test_export_is_not_for_viewers(app_client: AsyncClient, viewer_token, auth_headers):
    response = await app_client.get(
        "/api/blocked-domains/export", params={"range": "24h"}, headers=auth_headers(viewer_token)
    )
    assert response.status_code == 403


def _raw(domain: str, status: int) -> RawEvent:
    return RawEvent(
        timestamp=datetime.now(UTC), duration_ms=1, client_ip="10.0.0.1", action="TCP_DENIED",
        status_code=status, bytes=0, method="GET", url=f"http://{domain}/", domain=domain,
        user=None, blocked=True,
    )


async def test_events_block_kind(app_client: AsyncClient, admin_token, auth_headers, db_session):
    db_session.add_all([_raw("denied.example", 403), _raw("login.example", 407)])
    await db_session.commit()

    async def domains(**params):
        response = await app_client.get(
            "/api/events", params={"blocked_only": "true", **params}, headers=auth_headers(admin_token)
        )
        assert response.status_code == 200
        return sorted(item["domain"] for item in response.json()["items"])

    assert await domains() == ["denied.example", "login.example"]
    assert await domains(block_kind="denied") == ["denied.example"]
    assert await domains(block_kind="auth") == ["login.example"]
