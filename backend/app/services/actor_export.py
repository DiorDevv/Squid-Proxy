"""The "Who" leaderboard as a downloadable file -- every row, not one page.

Built in memory: the list is one row per actor (a few thousand at most, see
squid_ops_service.ACTOR_EXPORT_MAX_ROWS), unlike the raw-event exports in
export_service.py, which have to stream.
"""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.branch_scope import BranchScope, scope_label
from app.schemas.squid_ops import ActorRow
from app.services.export_service import _escape_csv_formula, _sanitize_xlsx_value

_MB = 1024 * 1024

# Column headers / sheet labels per UI language -- the file is read by
# people, so it follows the language they were using the dashboard in.
_TEXT = {
    "uz": {
        "headers": [
            "№", "Foydalanuvchi / IP", "Turi", "So'rovlar", "Bloklangan", "Bloklash %",
            "Yuklab olingan, MB", "Yuborilgan, MB", "Top kategoriya",
        ],
        "user": "Login", "ip": "IP", "sheet": "Kim", "info": "Ma'lumot",
        "branch": "Filial", "all_branches": "Barcha filiallar", "range": "Oraliq (UTC)",
        "kind": "Filtr", "search": "Qidiruv", "sort": "Saralash", "rows": "Qatorlar",
        "created": "Yaratilgan (UTC)",
        "kinds": {"all": "Hammasi", "user": "Faqat login", "ip": "Faqat IP"},
    },
    "en": {
        "headers": [
            "#", "User / IP", "Type", "Requests", "Blocked", "Blocked %",
            "Downloaded, MB", "Uploaded, MB", "Top category",
        ],
        "user": "Login", "ip": "IP", "sheet": "Who", "info": "Info",
        "branch": "Branch", "all_branches": "All branches", "range": "Range (UTC)",
        "kind": "Filter", "search": "Search", "sort": "Sort", "rows": "Rows",
        "created": "Created (UTC)",
        "kinds": {"all": "All", "user": "Users only", "ip": "IPs only"},
    },
    "ru": {
        "headers": [
            "№", "Пользователь / IP", "Тип", "Запросы", "Заблокировано", "Блокировки %",
            "Скачано, МБ", "Отправлено, МБ", "Топ категория",
        ],
        "user": "Логин", "ip": "IP", "sheet": "Кто", "info": "Сведения",
        "branch": "Филиал", "all_branches": "Все филиалы", "range": "Период (UTC)",
        "kind": "Фильтр", "search": "Поиск", "sort": "Сортировка", "rows": "Строк",
        "created": "Создано (UTC)",
        "kinds": {"all": "Все", "user": "Только логины", "ip": "Только IP"},
    },
}
LANGUAGES = tuple(_TEXT)


def _values(index: int, row: ActorRow, text: dict) -> list[object]:
    return [
        index,
        _escape_csv_formula(row.actor),
        text["user"] if row.is_user else text["ip"],
        row.request_count,
        row.blocked_count,
        round(row.blocked_ratio * 100, 1),
        round(row.total_bytes / _MB, 2),
        round(row.bytes_received / _MB, 2),
        row.top_category.value if row.top_category else "",
    ]


def actors_csv(rows: list[ActorRow], lang: str) -> bytes:
    text = _TEXT.get(lang, _TEXT["uz"])
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(text["headers"])
    for i, row in enumerate(rows, start=1):
        writer.writerow(_values(i, row, text))
    # BOM so Excel opens the UTF-8 (Cyrillic, o'/g') text correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


def actors_xlsx(
    rows: list[ActorRow],
    lang: str,
    *,
    branch: BranchScope,
    since: datetime,
    until: datetime,
    kind: str,
    search: str | None,
    sort: str,
) -> bytes:
    text = _TEXT.get(lang, _TEXT["uz"])
    wb = Workbook()
    ws = wb.active
    ws.title = text["sheet"]
    ws.append(text["headers"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=1):
        ws.append([_sanitize_xlsx_value(v) for v in _values(i, row, text)])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in enumerate([7, 34, 8, 12, 12, 11, 18, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    info = wb.create_sheet(text["info"])
    for label, value in [
        (text["branch"], scope_label(branch) or text["all_branches"]),
        (text["range"], f"{since:%Y-%m-%d %H:%M} — {until:%Y-%m-%d %H:%M}"),
        (text["kind"], text["kinds"].get(kind, kind)),
        (text["search"], _escape_csv_formula(search or "—")),
        (text["sort"], sort),
        (text["rows"], len(rows)),
        (text["created"], f"{datetime.now(until.tzinfo):%Y-%m-%d %H:%M}"),
    ]:
        info.append([label, _sanitize_xlsx_value(value)])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 40

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
