"""The "Who" list beyond its first page: offset paging, the All / Users /
IPs filter with its counts, and the full-list Excel/CSV export."""

import io
from datetime import UTC, datetime, timedelta

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit_log import AuditAction, AuditLogEntry
from app.models.client_aggregate import ClientMinuteAggregate
from app.services import squid_ops_service

NOW = datetime.now(UTC).replace(second=0, microsecond=0)
SINCE = NOW - timedelta(hours=1)
BUCKET = SINCE + timedelta(minutes=1)


async def _seed(db_session: AsyncSession) -> None:
    rows = [
        # 3 users with descending traffic...
        ClientMinuteAggregate(bucket_ts=BUCKET, client_ip=f"10.0.1.{i}", branch="main",
                              user=f"user{i}", request_count=100 - i)
        for i in range(3)
    ] + [
        # ...and 4 login-less IPs, two of them tied on requests
        ClientMinuteAggregate(bucket_ts=BUCKET, client_ip=f"10.0.2.{i}", branch="main",
                              user=None, request_count=50 if i < 2 else 10 - i)
        for i in range(4)
    ]
    db_session.add_all(rows)
    await db_session.commit()


async def _page(db, *, limit, offset=0, kind="all", search=None):
    return await squid_ops_service.get_actor_leaderboard(
        db, SINCE, NOW, branch=None, limit=limit, sort="requests", search=search, offset=offset, kind=kind
    )


async def test_pages_cover_every_actor_once_in_order(db_session: AsyncSession):
    await _seed(db_session)
    seen = []
    for offset in range(0, 7, 3):
        page = await _page(db_session, limit=3, offset=offset)
        assert (page.total, page.user_count, page.ip_count) == (7, 3, 4)
        assert page.actor_kind == "mixed"  # same on every page, even an all-IP one
        seen += [r.actor for r in page.rows]
    # ties (10.0.2.0 / 10.0.2.1 at 50) break by actor, so no row repeats or
    # goes missing across page boundaries
    assert seen == ["user0", "user1", "user2", "10.0.2.0", "10.0.2.1", "10.0.2.2", "10.0.2.3"]


async def test_kind_filter_and_its_counts(db_session: AsyncSession):
    await _seed(db_session)
    users = await _page(db_session, limit=50, kind="user")
    assert [r.actor for r in users.rows] == ["user0", "user1", "user2"]
    assert (users.total, users.user_count, users.ip_count, users.actor_kind) == (3, 3, 4, "user")

    ips = await _page(db_session, limit=50, kind="ip")
    assert all(not r.is_user for r in ips.rows) and len(ips.rows) == 4
    assert (ips.total, ips.actor_kind) == (4, "client_ip")

    # counts follow the search, so the tab badges match what's listed
    searched = await _page(db_session, limit=50, search="10.0.2")
    assert (searched.total, searched.user_count, searched.ip_count) == (4, 0, 4)


async def test_export_xlsx_has_every_row_and_is_audited(
    app_client: AsyncClient, admin_token, auth_headers, db_session: AsyncSession
):
    await _seed(db_session)
    response = await app_client.get(
        "/api/analytics/actors/export",
        params={"range": "1h", "format": "xlsx", "lang": "uz"},
        headers=auth_headers(admin_token),
    )
    assert response.status_code == 200
    assert 'filename="kim-barcha-' in response.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(response.content))["Kim"]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0][1] == "Foydalanuvchi / IP"
    assert [v[1] for v in values[1:]][:3] == ["user0", "user1", "user2"]
    assert len(values) == 1 + 7
    assert {v[2] for v in values[1:]} == {"Login", "IP"}

    entries = (
        await db_session.execute(select(AuditLogEntry).where(AuditLogEntry.action == AuditAction.EXPORT_CREATED))
    ).scalars().all()
    assert [e.detail for e in entries] == ["who-list format=xlsx kind=all rows=7"]


async def test_export_csv_respects_kind_and_neutralises_formulas(
    app_client: AsyncClient, admin_token, auth_headers, db_session: AsyncSession
):
    db_session.add(ClientMinuteAggregate(bucket_ts=BUCKET, client_ip="10.0.3.1", branch="main",
                                         user="=HYPERLINK(1)", request_count=5))
    await _seed(db_session)
    response = await app_client.get(
        "/api/analytics/actors/export",
        params={"range": "1h", "format": "csv", "kind": "user", "lang": "en"},
        headers=auth_headers(admin_token),
    )
    assert response.status_code == 200
    body = response.content.decode("utf-8-sig").splitlines()
    assert body[0].startswith("#,User / IP,Type")
    assert len(body) == 1 + 4  # users only
    assert any(line.split(",")[1] == "'=HYPERLINK(1)" for line in body[1:])


async def test_export_is_not_for_viewers(app_client: AsyncClient, viewer_token, auth_headers):
    response = await app_client.get(
        "/api/analytics/actors/export", params={"range": "1h"}, headers=auth_headers(viewer_token)
    )
    assert response.status_code == 403
